"""Persist experiment results and write the workbook template.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
import math
from pathlib import Path


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    cfg = result.get("cfg", {})
    exp_id = str(cfg.get("exp_id", "")).strip()
    if not exp_id or Path(exp_id).name != exp_id:
        raise ValueError("result.cfg.exp_id must be a nonempty filename stem")
    directory = Path(results_dir)
    directory.mkdir(parents=True, exist_ok=True)

    def json_safe(value):
        if isinstance(value, dict):
            return {str(k): json_safe(v) for k, v in value.items() if k != "best_state"}
        if isinstance(value, (list, tuple)):
            return [json_safe(v) for v in value]
        if hasattr(value, "tolist"):
            return json_safe(value.tolist())
        if hasattr(value, "item"):
            value = value.item()
        if isinstance(value, float) and not math.isfinite(value):
            return None
        return value

    payload = {key: result.get(key, {}) for key in ("cfg", "history", "summary")}
    path = directory / f"{exp_id}.json"
    if path.exists():
        raise FileExistsError(f"experiment result already exists; use a new exp_id: {path}")
    path.write_text(json.dumps(json_safe(payload), ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    directory = Path(results_dir)
    if not directory.exists():
        return []
    results = [json.loads(path.read_text(encoding="utf-8")) for path in directory.glob("*.json")]
    return sorted(results, key=lambda item: item.get("cfg", {}).get("exp_id", ""))


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg = result.get("cfg", {})
    summary = result.get("summary", {})
    optimizer_names = {
        "sgd": "SGD",
        "sgd_momentum": "SGD+momentum",
        "adam": "Adam",
        "adamw": "AdamW",
    }
    row = {
        "exp_id": cfg.get("exp_id"),
        "group": cfg.get("group"),
        "description": cfg.get("description"),
        "loss": {"ce": "CE", "mse": "MSE"}.get(cfg.get("loss"), cfg.get("loss")),
        "optimizer": optimizer_names.get(cfg.get("optimizer"), cfg.get("optimizer")),
        "lr": cfg.get("lr"),
        "weight_decay": cfg.get("weight_decay"),
        "batch": cfg.get("batch"),
        "epochs": cfg.get("epochs"),
        "hidden": "-".join(map(str, cfg.get("hidden", ()))) if cfg.get("hidden") is not None else None,
        "dropout": cfg.get("dropout"),
        "clip_norm": "none" if cfg.get("clip_norm") is None else cfg.get("clip_norm"),
        "precision": cfg.get("precision"),
        "init": cfg.get("init"),
        "seed": cfg.get("seed"),
        "step0_loss": summary.get("step0_loss"),
        "best_val_loss": summary.get("best_val_loss"),
        "best_epoch": summary.get("best_epoch"),
        "final_train_loss": summary.get("final_train_loss"),
        "final_val_loss": summary.get("final_val_loss"),
        "val_acc": summary.get("val_acc"),
        "val_macro_f1": summary.get("val_macro_f1"),
        "time_per_epoch_s": summary.get("time_per_epoch_s"),
        "peak_mem_MB": summary.get("peak_mem_MB"),
        "diverged": summary.get("diverged"),
        "eval_acc": None if eval_scores is None else eval_scores.get("accuracy", eval_scores.get("eval_acc")),
        "eval_macro_f1": None if eval_scores is None else eval_scores.get("macro_f1", eval_scores.get("eval_macro_f1")),
        "figure_file": f"figures/{cfg.get('exp_id', '')}.png",
        "notes": notes,
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               baseline_seed_ids: list[str] | None = None,
               summary_notes: dict[str, str] | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức (step0_gap_vs_lnC, gap_val_minus_train,
         delta_val_f1_vs_base, beyond_noise)
      4. wb.save(out_path)
    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.
    """
    from openpyxl import load_workbook

    workbook_path = Path(template_path)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"template workbook not found: {workbook_path}")
    wb = load_workbook(workbook_path)
    if "Experiments" not in wb.sheetnames:
        raise KeyError("template workbook must contain an 'Experiments' sheet")
    ws = wb["Experiments"]
    headers = {cell.value: cell.column for cell in ws[1] if cell.value is not None}
    formula_columns = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}
    if len(rows) > ws.max_row - 1:
        raise ValueError(f"template has room for at most {ws.max_row - 1} experiment rows")
    for row_index, row in enumerate(rows, start=2):
        for key, value in row.items():
            column = headers.get(key)
            if column is None or key in formula_columns:
                continue
            cell = ws.cell(row=row_index, column=column)
            if cell.data_type == "f":
                continue
            cell.value = value

    if "Seeds" not in wb.sheetnames or "Summary" not in wb.sheetnames:
        raise KeyError("template workbook must contain 'Seeds' and 'Summary' sheets")
    seed_ws = wb["Seeds"]
    seed_ids = list(baseline_seed_ids or [])
    summary_row = next(
        (row for row in range(2, seed_ws.max_row + 1) if seed_ws.cell(row=row, column=1).value == "mean"),
        seed_ws.max_row + 1,
    )
    seed_rows = list(range(2, max(2, summary_row - 1)))  # leave the spacer and summary rows intact
    if len(seed_ids) > len(seed_rows):
        raise ValueError(f"Seeds sheet has room for at most {len(seed_rows)} baseline runs")
    for offset, row_index in enumerate(seed_rows):
        seed_ws.cell(row=row_index, column=1).value = seed_ids[offset] if offset < len(seed_ids) else None

    summary_ws = wb["Summary"]
    group_notes = summary_notes or {}
    for row_index in range(2, summary_ws.max_row + 1):
        group = summary_ws.cell(row=row_index, column=1).value
        if group in group_notes:
            summary_ws.cell(row=row_index, column=8).value = group_notes[group]

    # Keep Excel as the calculation engine for the template's existing formulas.
    if getattr(wb, "calculation", None) is not None:
        wb.calculation.calcMode = "auto"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True

    output = Path(out_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
